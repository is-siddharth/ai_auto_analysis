"""Release-level smoke test for the AI Data Analytics Assistant.

Run from the project root:
    python scripts/verify_build.py

This checks both import-level integration and, when the bundled Superstore
sample is available, the complete workbook generation path.
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from src.analysis.charts import recommend_charts
from src.analysis.engine import run_analysis
from src.analysis.excel_functions import generate_excel_function_recommendations
from src.analysis.filters import recommend_filters
from src.analysis.kpis import calculate_kpi, generate_kpi_recommendations
from src.analysis.recommendations import generate_analysis_questions
from src.dashboard.generator import recommend_dashboard_layouts
from src.excel.workbook import generate_excel_workbook
from src.ingestion.loader import load_raw_dataset
from src.pipeline import analyze_dataset


def main():
    checks = {
        "pipeline": analyze_dataset,
        "analysis recommendations": generate_analysis_questions,
        "KPI recommendations": generate_kpi_recommendations,
        "KPI calculation": calculate_kpi,
        "Excel function recommendations": generate_excel_function_recommendations,
        "chart recommendations": recommend_charts,
        "filter recommendations": recommend_filters,
        "dashboard layout recommendations": recommend_dashboard_layouts,
        "workbook generation": generate_excel_workbook,
        "analysis engine": run_analysis,
    }

    print("AI Data Analytics Assistant - release verification")
    print("=" * 56)
    for name, obj in checks.items():
        if not callable(obj):
            raise RuntimeError(f"Build check failed: {name} is not callable.")
        print(f"[OK] {name}")

    sample = Path(__file__).resolve().parents[1] / "data" / "raw" / "Superstore.csv"
    if sample.exists():
        print("\nRunning end-to-end workbook smoke test...")
        result = analyze_dataset(str(sample))
        df = result["data"]
        raw_df = load_raw_dataset(str(sample))

        selected_kpis = []
        for name in ["Total Orders", "Total Sales", "Profit Margin"]:
            matches = result["kpi_recommendations"].loc[result["kpi_recommendations"]["KPI"] == name]
            if not matches.empty:
                selected_kpis.append(calculate_kpi(df, matches.iloc[0]))

        selected_analyses = []
        for index, recommendation in result["analysis_questions"].iterrows():
            if recommendation["Analysis Type"] in {"Dimension Comparison", "Ranking", "Time Trend"}:
                chart_match = result["chart_recommendations"].loc[
                    result["chart_recommendations"]["Analysis Index"] == index
                ]
                chart_label = chart_match.iloc[0]["Recommended Chart"] if not chart_match.empty else "Table"
                selected_analyses.append({
                    "recommendation": recommendation.to_dict(),
                    "result": run_analysis(df, recommendation),
                    "chart_label": chart_label,
                })
            if len(selected_analyses) == 3:
                break

        filter_recommendations = recommend_filters(df, result["analytical_fields"])
        selected_filters = filter_recommendations["Field"].tolist()[:3] if not filter_recommendations.empty else []

        layouts = recommend_dashboard_layouts(
            result["analytical_fields"], len(selected_kpis), len(selected_analyses)
        )
        if len(layouts) != 5:
            raise RuntimeError("Dashboard recommender did not return five layouts.")

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as temp_file:
            output = Path(temp_file.name)

        try:
            generate_excel_workbook(
                output_path=output,
                raw_df=raw_df,
                cleaned_df=df,
                data_dictionary=result["data_dictionary"],
                kpis=selected_kpis,
                analyses=selected_analyses,
                excel_function_recommendations=result["excel_function_recommendations"],
                dashboard_layout=layouts[0],
                dataset_name=sample.name,
                selected_filters=selected_filters,
                filter_recommendations=filter_recommendations,
            )
            wb = load_workbook(output)
            if wb.sheetnames[-1] != "Dashboard_Template":
                raise RuntimeError("Dashboard_Template is not the final worksheet.")
            if not any(name.startswith("KPI_") for name in wb.sheetnames):
                raise RuntimeError("No KPI worksheet was generated.")
            if not any(name.startswith("Analysis_") for name in wb.sheetnames):
                raise RuntimeError("No analysis worksheet was generated.")
            if "Filter_Guide" not in wb.sheetnames:
                raise RuntimeError("Filter_Guide worksheet was not generated.")
            print(f"[OK] end-to-end workbook smoke test ({len(wb.sheetnames)} sheets)")
        finally:
            output.unlink(missing_ok=True)

    else:
        print("[SKIP] bundled Superstore sample not found; import checks still passed.")

    print("\nRelease verification passed.")


if __name__ == "__main__":
    main()
