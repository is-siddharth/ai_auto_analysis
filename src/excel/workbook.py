"""Final Excel workbook generator for the AI Data Analytics Assistant."""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
GRAY = "F3F4F6"
DARK_GRAY = "4B5563"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"
BORDER = "D1D5DB"

THIN = Side(style="thin", color=BORDER)
BOX_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _safe_sheet_title(title: str, used: set[str]) -> str:
    title = re.sub(r"[\\/*?:\[\]]", "-", str(title)).strip() or "Sheet"
    title = title[:31]
    base = title
    counter = 2
    while title in used:
        suffix = f"_{counter}"
        title = f"{base[:31-len(suffix)]}{suffix}"
        counter += 1
    used.add(title)
    return title


def _style_title(ws, title: str, subtitle: str | None = None, end_col: int = 8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(size=18, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        cell = ws.cell(2, 1, subtitle)
        cell.font = Font(size=10, italic=True, color=DARK_GRAY)
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 32


def _style_header_row(ws, row: int, start_col: int, end_col: int):
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX_BORDER


def _autofit(ws, min_width=10, max_width=42):
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, max((len(line) for line in value.split("\n")), default=0))
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, min_width), max_width)


def _excel_safe_value(value):
    """Convert pandas/numpy missing values into Excel-safe Python values."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        try:
            value = value.item()
        except Exception:
            pass
    return value


def _write_dataframe(ws, df: pd.DataFrame, start_row: int = 4, start_col: int = 1, table_name: str | None = None):
    if df is None:
        df = pd.DataFrame()
    df = df.copy()
    rows = list(dataframe_to_rows(df, index=False, header=True))
    for row_index, row in enumerate(rows, start=start_row):
        for col_index, value in enumerate(row, start=start_col):
            cell = ws.cell(row_index, col_index, _excel_safe_value(value))
            cell.border = BOX_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (pd.Timestamp,)):
                cell.value = cell.value.to_pydatetime()
                cell.number_format = "dd-mmm-yyyy"
    if len(df.columns) and len(df) >= 1 and table_name:
        ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(start_col + len(df.columns) - 1)}{start_row + len(df)}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    if len(df.columns):
        _style_header_row(ws, start_row, start_col, start_col + len(df.columns) - 1)
    ws.freeze_panes = ws.cell(start_row + 1, start_col)
    _autofit(ws)
    return start_row + len(df) + 1

def _clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Return the approved working dataset without silently modifying it.

    Cleaning decisions are made upstream through the approval-based cleaning
    workflow. The workbook layer must not perform additional hidden changes.
    """
    return df.copy()


def _write_kpi_sheet(wb, used, kpi: dict[str, Any], position: int):
    title = _safe_sheet_title(f"KPI_{position:02d}_{kpi['KPI']}", used)
    ws = wb.create_sheet(title)
    _style_title(ws, kpi["KPI"], "Selected KPI — calculated directly from the uploaded dataset.", 6)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A4:F7")
    value_cell = ws["A4"]
    value = kpi.get("Value")
    if value is None:
        display = "N/A"
    elif kpi.get("Category") == "Ratio":
        display = value
        value_cell.number_format = "0.0%"
    elif isinstance(value, float):
        display = value
        value_cell.number_format = "#,##0.00"
    else:
        display = value
        value_cell.number_format = "#,##0"
    value_cell.value = display
    value_cell.font = Font(size=28, bold=True, color=NAVY)
    value_cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell.border = BOX_BORDER

    details = pd.DataFrame([
        ["Category", kpi.get("Category", "")],
        ["Calculation", kpi.get("Calculation", "")],
        ["Source Fields", kpi.get("Fields", "")],
        ["Formula Logic", kpi.get("Formula Logic", "")],
        ["Why It Matters", kpi.get("Rationale", "")],
    ], columns=["Item", "Detail"])
    _write_dataframe(ws, details, start_row=9, table_name=f"KPIInfo{position:02d}")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 18
    return title


def _add_excel_chart(ws, result: dict[str, Any], anchor: str = "A1"):
    chart_type = result.get("chart_type")
    table = result.get("table")
    x = result.get("x")
    y = result.get("y")
    if chart_type not in {"bar", "line"} or not isinstance(table, pd.DataFrame) or table.empty:
        return False
    if x not in table.columns or y not in table.columns:
        return False

    # Keep chart source data far outside the visible area. Never hide visible chart/table cells.
    start_col = 40
    source_row = 2
    ws.cell(source_row, start_col, x)
    ws.cell(source_row, start_col + 1, y)
    visible_rows = table[[x, y]].head(30)
    for idx, (_, row) in enumerate(visible_rows.iterrows(), start=source_row + 1):
        ws.cell(idx, start_col, _excel_safe_value(row[x]))
        ws.cell(idx, start_col + 1, _excel_safe_value(row[y]))
    ws.column_dimensions[get_column_letter(start_col)].hidden = True
    ws.column_dimensions[get_column_letter(start_col + 1)].hidden = True

    data = Reference(ws, min_col=start_col + 1, min_row=source_row, max_row=source_row + len(visible_rows))
    cats = Reference(ws, min_col=start_col, min_row=source_row + 1, max_row=source_row + len(visible_rows))
    chart = LineChart() if chart_type == "line" else BarChart()
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.title = f"{y} by {x}"
    chart.y_axis.title = y
    chart.x_axis.title = x
    chart.height = 8.5
    chart.width = 14.5
    chart.style = 13 if chart_type == "line" else 10
    if chart_type == "bar":
        chart.type = "bar"
    ws.add_chart(chart, anchor)
    return True

def _write_analysis_sheet(wb, used, recommendation: dict[str, Any], result: dict[str, Any], position: int, chart_label: str):
    analysis_type = str(recommendation.get("Analysis Type", "Analysis"))
    title = _safe_sheet_title(f"Analysis_{position:02d}", used)
    ws = wb.create_sheet(title)
    _style_title(ws, f"Analysis {position:02d}", recommendation.get("Question", ""), 10)
    ws.sheet_view.showGridLines = False

    meta = pd.DataFrame([
        ["Analysis Type", analysis_type],
        ["Priority", recommendation.get("Priority", "")],
        ["Recommended Visual", chart_label],
        ["Rationale", recommendation.get("Rationale", "")],
    ], columns=["Item", "Detail"])
    meta_end = _write_dataframe(ws, meta, start_row=4, table_name=f"AnalysisMeta{position:02d}")

    table = result.get("table") if result else pd.DataFrame()
    start_row = max(11, meta_end + 2)
    chart_anchor_row = start_row
    if isinstance(table, pd.DataFrame) and not table.empty:
        ws.cell(start_row - 1, 1, "Analysis Summary Table")
        ws.cell(start_row - 1, 1).font = Font(size=13, bold=True, color=NAVY)
        table_end = _write_dataframe(ws, table, start_row=start_row, table_name=f"AnalysisData{position:02d}")
        chart_anchor_row = table_end + 3
    else:
        ws.cell(start_row, 1, "No tabular result was produced for this analysis.")
        ws.cell(start_row, 1).font = Font(italic=True, color=DARK_GRAY)
        chart_anchor_row = start_row + 2

    ws.cell(chart_anchor_row - 1, 1, "Generated Visual")
    ws.cell(chart_anchor_row - 1, 1).font = Font(size=13, bold=True, color=NAVY)
    _add_excel_chart(ws, result or {}, anchor=f"A{chart_anchor_row}")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 62
    ws.sheet_view.zoomScale = 85
    return title

def _write_excel_guide(wb, used, recommendations: pd.DataFrame):
    title = _safe_sheet_title("Excel_Guide", used)
    ws = wb.create_sheet(title)
    _style_title(ws, "Excel Learning Guide", "Dataset-specific Excel functions suggested by the analytical opportunities detected in this dataset.", 9)
    if recommendations is None or recommendations.empty:
        ws["A4"] = "No specific Excel-function opportunities were detected."
        return title
    _write_dataframe(ws, recommendations, start_row=4, table_name="ExcelLearningGuide")
    return title


def _layout_cells(layout_key: str):
    layouts = {
        "kpi_left_slicer": {"kpis": ["B5:E8", "F5:I8", "J5:M8", "N5:Q8"], "slicer": "B10:C27", "charts": ["D10:J18", "K10:Q18", "D20:J28", "K20:Q28"]},
        "kpi_top_trend": {"kpis": ["B5:E8", "F5:I8", "J5:M8", "N5:Q8"], "slicer": None, "hero": "B10:Q21", "charts": ["B23:I31", "J23:Q31"]},
        "hero_chart": {"kpis": ["B5:E8", "F5:I8", "J5:M8"], "slicer": None, "hero": "B10:Q23", "charts": ["B25:I33", "J25:Q33"]},
        "comparison_board": {"kpis": ["B5:E8", "F5:I8", "J5:M8"], "slicer": "B10:C28", "charts": ["D10:J18", "K10:Q18", "D20:J28", "K20:Q28"]},
        "balanced_grid": {"kpis": ["B5:E8", "F5:I8", "J5:M8", "N5:Q8"], "slicer": "B10:C20", "charts": ["D10:J18", "K10:Q18", "B22:I30", "J22:Q30"]},
    }
    return layouts.get(layout_key, layouts["balanced_grid"])


def _fill_placeholder(ws, cell_range: str, label: str, detail: str, fill_color: str = PALE_BLUE):
    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    ws.merge_cells(cell_range)
    cell = ws.cell(min_row, min_col, f"{label}\n{detail}")
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.border = BOX_BORDER
    cell.font = Font(size=11, bold=True, color="111827")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for item in row:
            item.border = BOX_BORDER
    return cell


def _write_dashboard_template(wb, used, layout, kpi_sheets, analysis_sheets, analysis_labels, dataset_name, selected_filters):
    title = _safe_sheet_title("Dashboard_Template", used)
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B5"
    for col in range(1, 18):
        ws.column_dimensions[get_column_letter(col)].width = 12
    for row in range(1, 40):
        ws.row_dimensions[row].height = 24
    ws.merge_cells("A1:Q2")
    ws["A1"] = f"DASHBOARD TEMPLATE — {dataset_name}"
    ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A3:Q3")
    ws["A3"] = f"Selected layout: {layout['Layout']} | Use the placeholders as a placement guide. Selected filters: {', '.join(selected_filters) if selected_filters else 'None'}"
    ws["A3"].font = Font(size=10, color="111827", italic=True)
    ws["A3"].fill = PatternFill("solid", fgColor="E2E8F0")
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="center")

    cells = _layout_cells(layout["Key"])
    for pos, cell_range in enumerate(cells.get("kpis", []), start=1):
        source = kpi_sheets[pos - 1] if pos <= len(kpi_sheets) else "Open space / optional"
        _fill_placeholder(ws, cell_range, f"KPI {pos}", source, "DBEAFE")

    if cells.get("slicer"):
        filter_text = "\n".join(selected_filters) if selected_filters else "No filters selected"
        _fill_placeholder(ws, cells["slicer"], "FILTER / SLICER AREA", filter_text, "FEF3C7")

    chart_ranges = []
    if cells.get("hero"):
        chart_ranges.append(cells["hero"])
    chart_ranges.extend(cells.get("charts", []))
    for pos, cell_range in enumerate(chart_ranges, start=1):
        source = analysis_sheets[pos - 1] if pos <= len(analysis_sheets) else "Open space / optional"
        label = analysis_labels[pos - 1] if pos <= len(analysis_labels) else "Additional analysis"
        _fill_placeholder(ws, cell_range, f"CHART {pos}", f"{source}\n{label}", "FFFFFF")

    ws.merge_cells("A35:Q38")
    ws["A35"] = "BUILD NOTE\nThe workbook contains the calculated KPI values, analysis tables, and generated charts on their own sheets. Place those visuals into the marked areas. The selected filter fields are the recommended dashboard controls; create native Excel slicers manually from the relevant table/PivotTable if desired."
    ws["A35"].fill = PatternFill("solid", fgColor="E2E8F0")
    ws["A35"].font = Font(color="111827", italic=True)
    ws["A35"].alignment = Alignment(wrap_text=True, vertical="center")
    return title


def _write_dashboard_instructions(wb, used, layout, kpi_sheets, analysis_sheets, analysis_labels, selected_filters):
    title = _safe_sheet_title("Dashboard_Instructions", used)
    ws = wb.create_sheet(title)
    _style_title(ws, "Dashboard Build Instructions", "Follow this mapping to assemble the selected dashboard layout in Excel.", 8)
    rows = []
    for idx, sheet in enumerate(kpi_sheets, start=1):
        rows.append({"Component": f"KPI {idx}", "Source Sheet": sheet, "Place In": f"KPI {idx} placeholder", "What To Use": "The large KPI value and label", "Notes": "Keep the KPI label visible."})
    for idx, sheet in enumerate(analysis_sheets, start=1):
        label = analysis_labels[idx - 1] if idx - 1 < len(analysis_labels) else "Selected analysis"
        rows.append({"Component": f"Chart {idx}", "Source Sheet": sheet, "Place In": f"Chart {idx} placeholder", "What To Use": "The generated chart; keep its supporting table traceable", "Notes": label})
    for idx, field in enumerate(selected_filters, start=1):
        rows.append({"Component": f"Filter {idx}", "Source Sheet": "02_Cleaned_Data", "Place In": "Filter / Slicer Area", "What To Use": field, "Notes": "Create an Excel slicer from a table/PivotTable field, or use the table's filter dropdown."})
    if not rows:
        rows.append({"Component": "No selected outputs", "Source Sheet": "-", "Place In": "-", "What To Use": "Return to the Analysis tab", "Notes": "Select KPIs or analyses before generating."})
    _write_dataframe(ws, pd.DataFrame(rows), start_row=4, table_name="DashboardInstructions")
    return title


def _write_filter_guide(wb, used, filter_recommendations, selected_filters):
    title = _safe_sheet_title("Filter_Guide", used)
    ws = wb.create_sheet(title)
    _style_title(ws, "Dashboard Filter Guide", "What each recommended filter does and why it can be useful.", 8)
    if filter_recommendations is None or filter_recommendations.empty:
        ws["A4"] = "No filter recommendations were generated for this dataset."
        return title
    guide = filter_recommendations.copy()
    guide["Selected"] = guide["Field"].isin(selected_filters)
    _write_dataframe(ws, guide, start_row=4, table_name="FilterGuide")
    return title


def generate_excel_workbook(
    output_path: str | os.PathLike,
    raw_df: pd.DataFrame,
    cleaned_df: pd.DataFrame,
    data_dictionary: pd.DataFrame,
    kpis: list[dict[str, Any]],
    analyses: list[dict[str, Any]],
    excel_function_recommendations: pd.DataFrame,
    dashboard_layout: dict[str, Any],
    dataset_name: str,
    selected_filters: list[str] | None = None,
    filter_recommendations: pd.DataFrame | None = None,
    cleaning_audit: pd.DataFrame | None = None,
) -> str:
    """Generate the final analyst workbook and return its path."""
    selected_filters = selected_filters or []
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    used: set[str] = set()

    # Read-me / navigation first.
    readme_title = _safe_sheet_title("00_Read_Me", used)
    ws = wb.create_sheet(readme_title)
    _style_title(ws, "AI Data Analytics Assistant — Workbook", "Generated analytical foundation. Review the source sheets, then build the dashboard using the final template.", 8)
    overview = pd.DataFrame([
        ["Dataset", dataset_name],
        ["Rows", len(raw_df)],
        ["Columns", len(raw_df.columns)],
        ["Selected KPIs", len(kpis)],
        ["Selected Analyses", len(analyses)],
        ["Dashboard Layout", dashboard_layout.get("Layout", "")],
        ["Important", "Analysis sheets contain calculated summary tables and generated charts. Dashboard assembly is intentionally left for the user."],
    ], columns=["Item", "Value"])
    _write_dataframe(ws, overview, start_row=4, table_name="WorkbookOverview")

    raw_title = _safe_sheet_title("01_Raw_Data", used)
    ws = wb.create_sheet(raw_title)
    _style_title(ws, "Raw Data", "Original uploaded values. No analytical calculations are performed on this sheet.", min(12, len(raw_df.columns)))
    _write_dataframe(ws, raw_df, start_row=4, table_name="RawData")

    clean_title = _safe_sheet_title("02_Cleaned_Data", used)
    cleaned_for_output = _clean_dataframe(cleaned_df)
    ws = wb.create_sheet(clean_title)
    _style_title(ws, "Cleaned Data", "Approved working dataset. No additional cleaning is performed during workbook generation.", min(12, len(cleaned_for_output.columns)))
    _write_dataframe(ws, cleaned_for_output, start_row=4, table_name="CleanedData")

    dictionary_title = _safe_sheet_title("03_Data_Dictionary", used)
    ws = wb.create_sheet(dictionary_title)
    _style_title(ws, "Data Dictionary", "Detected analytical role and deterministic description for every field.", 8)
    _write_dataframe(ws, data_dictionary, start_row=4, table_name="DataDictionary")

    cleaning_title = _safe_sheet_title("04_Cleaning_Log", used)
    ws = wb.create_sheet(cleaning_title)
    _style_title(
        ws,
        "Cleaning Log",
        "Only approved cleaning actions are recorded here. An empty log means no cleaning actions were applied.",
        8,
    )
    if cleaning_audit is None or cleaning_audit.empty:
        ws["A4"] = "No cleaning actions were applied to the working dataset."
    else:
        _write_dataframe(ws, cleaning_audit, start_row=4, table_name="CleaningLog")

    kpi_sheet_names = []
    for position, kpi in enumerate(kpis, start=1):
        kpi_sheet_names.append(_write_kpi_sheet(wb, used, kpi, position))

    analysis_sheet_names = []
    analysis_labels = []
    for position, analysis in enumerate(analyses, start=1):
        analysis_sheet_names.append(_write_analysis_sheet(wb, used, analysis["recommendation"], analysis["result"], position, analysis.get("chart_label", "Table")))
        analysis_labels.append(str(analysis["recommendation"].get("Question", "Selected analysis")))

    _write_excel_guide(wb, used, excel_function_recommendations)
    _write_filter_guide(wb, used, filter_recommendations, selected_filters)
    _write_dashboard_instructions(wb, used, dashboard_layout, kpi_sheet_names, analysis_sheet_names, analysis_labels, selected_filters)
    _write_dashboard_template(wb, used, dashboard_layout, kpi_sheet_names, analysis_sheet_names, analysis_labels, dataset_name, selected_filters)

    # Workbook-level usability settings.
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.outlinePr.summaryBelow = True
        ws.sheet_view.zoomScale = 90

    output_path = str(Path(output_path))
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
