import os

import pandas as pd
from openpyxl import load_workbook

from src.dashboard.generator import recommend_dashboard_layouts
from src.analysis.filters import recommend_filters
from src.excel.workbook import generate_excel_workbook


def test_dashboard_recommender_returns_five_visual_layouts():
    layouts = recommend_dashboard_layouts(
        {"dates": ["Order Date"], "dimensions": ["Category"], "measures": ["Sales"]},
        selected_kpi_count=3,
        selected_analysis_count=3,
    )
    assert len(layouts) == 5
    assert all(layout["Layout"] for layout in layouts)
    assert all(layout["Preview"] for layout in layouts)


def test_filter_recommendations_are_dataset_aware():
    df = pd.DataFrame({"Region": ["East", "West", "East"], "Category": ["A", "B", "A"], "Customer": ["x", "y", "z"], "Sales": [1, 2, 3]})
    result = recommend_filters(df, {"dimensions": ["Region", "Category", "Customer"]})
    assert "Region" in result["Field"].tolist()
    assert "Category" in result["Field"].tolist()
    assert "Customer" not in result["Field"].tolist()


def test_workbook_contains_expected_structure_and_final_dashboard(tmp_path):
    raw = pd.DataFrame({
        "Order ID": ["A", "B", "C"],
        "Category": ["Tech", "Office", "Tech"],
        "Sales": [100.0, 200.0, 150.0],
        "Order Date": pd.to_datetime(["2026-01-01", "2026-02-01", "2026-03-01"]),
    })
    dictionary = pd.DataFrame({
        "Column": raw.columns,
        "Classification": ["Potential ID", "Categorical Dimension", "Numerical Measure", "Date / Time"],
        "Description": ["ID", "Dimension", "Measure", "Date"],
        "Data Type": ["object", "object", "float64", "datetime64[ns]"],
        "Unique Values": [3, 2, 3, 3],
        "Missing Values": [0, 0, 0, 0],
    })
    kpis = [{
        "KPI": "Total Sales",
        "Value": 450.0,
        "Category": "Measure",
        "Calculation": "Sum of Sales",
        "Fields": "Sales",
        "Formula Logic": "SUM(Sales)",
        "Rationale": "Total measure",
    }]
    analysis_result = {
        "table": pd.DataFrame({"Category": ["Tech", "Office"], "Total": [250.0, 200.0]}),
        "chart_type": "bar",
        "x": "Category",
        "y": "Total",
    }
    analyses = [{
        "recommendation": {
            "Question": "How does Sales vary across Category?",
            "Analysis Type": "Dimension Comparison",
            "Priority": "High",
            "Rationale": "Compare groups.",
        },
        "result": analysis_result,
        "chart_label": "Bar Chart",
    }]
    layout = recommend_dashboard_layouts(
        {"dates": ["Order Date"], "dimensions": ["Category"], "measures": ["Sales"]}, 1, 1
    )[0]

    output = tmp_path / "output.xlsx"
    generate_excel_workbook(
        output, raw, raw, dictionary, kpis, analyses,
        pd.DataFrame(columns=["Task", "Recommended Function"]),
        layout, "test.csv",
        selected_filters=["Category"],
        filter_recommendations=recommend_filters(raw, {"dimensions": ["Category"]}),
    )

    wb = load_workbook(output)
    assert wb.sheetnames[:4] == ["00_Read_Me", "01_Raw_Data", "02_Cleaned_Data", "03_Data_Dictionary"]
    assert any(name.startswith("KPI_01_") for name in wb.sheetnames)
    assert "Analysis_01" in wb.sheetnames
    assert "Dashboard_Instructions" in wb.sheetnames
    assert "Filter_Guide" in wb.sheetnames
    assert "04_Cleaning_Log" in wb.sheetnames
    assert wb.sheetnames[-1] == "Dashboard_Template"
    assert len(wb["Analysis_01"]._charts) == 1
    assert "KPI 1" in wb["Dashboard_Template"]["B5"].value
